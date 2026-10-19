import os
import csv
import json
from pathlib import Path
from typing import List, Callable, Optional
from pypdf import PdfReader
from docx import Document as DocxDocument

from src.core.models import Document, DocumentPage

class DocumentLoader:
    """Universal document loader supporting PDF, DOCX, TXT, CSV, JSON, and source files."""

    SUPPORTED_EXTENSIONS = {
        ".pdf", ".docx", ".txt", ".md", ".markdown",
        ".csv", ".tsv", ".json", ".py", ".js", ".ts",
        ".html", ".css", ".sql", ".sh", ".yaml", ".yml", ".xml"
    }

    def __init__(self, progress_callback: Optional[Callable[[str, str, dict], None]] = None):
        self.progress_callback = progress_callback

    def _notify(self, stage: str, message: str, data: Optional[dict] = None) -> None:
        if self.progress_callback:
            self.progress_callback(stage, message, data or {})

    def load(self, file_path: str) -> Document:
        """Loads and parses a document from a local file path with page-level tracking."""
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        if not path.is_file():
            raise ValueError(f"Path is not a regular file: {file_path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format: '{ext}'. "
                f"Supported extensions: {', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
            )

        self._notify("PARSE", f"Reading file: {path.name}", {"filename": path.name, "format": ext})

        if ext == ".pdf":
            pages = self._load_pdf(path)
        elif ext == ".docx":
            pages = self._load_docx(path)
        elif ext in {".csv", ".tsv"}:
            pages = self._load_csv(path, delimiter="," if ext == ".csv" else "\t")
        elif ext == ".json":
            pages = self._load_json(path)
        else:
            pages = self._load_text_file(path)

        total_pages = len(pages)
        self._notify(
            "PARSE",
            f"Successfully parsed {path.name} ({total_pages} page{'s' if total_pages != 1 else ''})",
            {"filename": path.name, "total_pages": total_pages}
        )

        return Document(
            source_path=str(path),
            filename=path.name,
            file_type=ext.lstrip("."),
            pages=pages,
            total_pages=total_pages,
            metadata={
                "file_size_bytes": path.stat().st_size,
                "file_extension": ext,
            }
        )

    def _load_pdf(self, path: Path) -> List[DocumentPage]:
        pages: List[DocumentPage] = []
        reader = PdfReader(str(path))
        num_pages = len(reader.pages)

        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                pages.append(DocumentPage(page_number=idx + 1, text=text))
            else:
                pages.append(DocumentPage(page_number=idx + 1, text="[No extractable text on this page]"))

        if not pages:
            pages.append(DocumentPage(page_number=1, text="[Document contains no pages]"))

        return pages

    def _load_docx(self, path: Path) -> List[DocumentPage]:
        doc = DocxDocument(str(path))
        paragraphs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([c.text.strip() for c in row.cells if c.text.strip()])
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            return [DocumentPage(page_number=1, text="[Empty document]")]

        # Partition into ~3000 character page estimates if single stream
        pages: List[DocumentPage] = []
        chunk_size = 3000
        lines = full_text.split("\n\n")
        current_page = []
        current_len = 0
        page_no = 1

        for line in lines:
            if current_len + len(line) > chunk_size and current_page:
                pages.append(DocumentPage(page_number=page_no, text="\n\n".join(current_page)))
                page_no += 1
                current_page = [line]
                current_len = len(line)
            else:
                current_page.append(line)
                current_len += len(line)

        if current_page:
            pages.append(DocumentPage(page_number=page_no, text="\n\n".join(current_page)))

        return pages

    def _load_csv(self, path: Path, delimiter: str = ",") -> List[DocumentPage]:
        lines = []
        with open(path, mode="r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f, delimiter=delimiter)
            header = next(reader, None)
            header_str = f"Columns: {', '.join(header)}" if header else ""
            if header_str:
                lines.append(header_str)

            for row_idx, row in enumerate(reader):
                cleaned_row = [col.strip() for col in row if col.strip()]
                if cleaned_row:
                    lines.append(f"Row {row_idx + 1}: " + " | ".join(cleaned_row))

        full_text = "\n".join(lines)
        if not full_text.strip():
            return [DocumentPage(page_number=1, text="[Empty CSV file]")]

        return [DocumentPage(page_number=1, text=full_text)]

    def _load_json(self, path: Path) -> List[DocumentPage]:
        with open(path, mode="r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)

        formatted_text = json.dumps(data, indent=2)
        return [DocumentPage(page_number=1, text=formatted_text)]

    def _load_text_file(self, path: Path) -> List[DocumentPage]:
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        content = ""
        for enc in encodings:
            try:
                with open(path, mode="r", encoding=enc) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if not content.strip():
            content = "[Empty document]"

        return [DocumentPage(page_number=1, text=content)]
